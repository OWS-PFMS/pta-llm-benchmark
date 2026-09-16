#!/usr/bin/env python3
"""Table 4 - factor-level shift in influence and dependence, expert -> LLM.

Emits two artifacts from one pass:
  * table4_shifts.md    - plain markdown, same house style as tables_1-3.md
  * table4_shifts.docx  - Word table whose delta cells are shaded with the
                          Figure 4 diverging ramp, so the table doubles as
                          the heatmap without needing the figure beside it.

Shifts, row order and colors are imported from the Figure 4 generator rather
than recomputed, so the table and the figure cannot drift apart.
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(OUT.parent / "figures"))
from make_figure4_shift_heatmap import (  # noqa: E402
    CMAP, COLS, SHORT, shifts)

VMAX = 0.55  # matches the Figure 4 colorbar clamp
MINUS = "−"

CAPTION = (
    "Factor-level shift in structural position between the expert and LLM "
    "codings of the same 232 transcripts. Each value is the LLM score minus "
    "the expert score on that axis, in normalized units (weighted out-degree "
    "for influence, weighted in-degree for dependence, each divided by the "
    "largest factor in its own network). Positive values mean the LLM scored "
    "the factor higher; negative values mean the expert coding did. Total "
    "displacement is the sum of the two systems' Euclidean shifts - the "
    "combined length of the factor's displacement arrows in Figure 3 - and "
    "orders the rows. Cell shading follows Figure 4.")


def fmt(v):
    return "0.00" if abs(v) < 0.005 else f"{v:+.2f}".replace("-", MINUS)


def shade(cell, rgba):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"),
           "".join(f"{int(round(c * 255)):02X}" for c in rgba[:3]))
    cell._tc.get_or_add_tcPr().append(el)


def is_dark(rgba):
    return 0.2126 * rgba[0] + 0.7152 * rgba[1] + 0.0722 * rgba[2] < 0.42


df = shifts()

# ---- markdown -------------------------------------------------------------
md = ["**Table 4.** " + CAPTION.replace(" - ", " — ") + "\n",
      "| Factor | Challenge: Δ Influence | Challenge: Δ Dependence "
      "| Solution: Δ Influence | Solution: Δ Dependence "
      "| Total displacement |",
      "|:---|---:|---:|---:|---:|---:|"]
for factor, r in df.iterrows():
    md.append("| " + " | ".join(
        [SHORT.get(factor, factor)]
        + [fmt(r[k]) for k, _ in COLS]
        + [f"{r['total_disp']:.2f}"]) + " |")
md.append("")
(OUT / "table4_shifts.md").write_text("\n".join(md), encoding="utf-8")

# ---- docx -----------------------------------------------------------------
doc = Document()
for m in ("left_margin", "right_margin"):
    setattr(doc.sections[0], m, Inches(0.9))

cap = doc.add_paragraph()
cap.add_run("Table 4. ").bold = True
cap.add_run(CAPTION.replace(" - ", " — "))

table = doc.add_table(rows=2, cols=6)
table.style = "Table Grid"
table.cell(0, 0).merge(table.cell(1, 0)).text = "Factor"
table.cell(0, 1).merge(table.cell(0, 2)).text = "Challenge system"
table.cell(0, 3).merge(table.cell(0, 4)).text = "Solution system"
table.cell(0, 5).merge(table.cell(1, 5)).text = "Total displacement"
for i, t in enumerate(["Δ Influence", "Δ Dependence"] * 2, start=1):
    table.cell(1, i).text = t

for factor, r in df.iterrows():
    cells = table.add_row().cells
    cells[0].text = SHORT.get(factor, factor)
    for i, (key, _) in enumerate(COLS, start=1):
        v = r[key]
        rgba = CMAP((v + VMAX) / (2 * VMAX))
        cells[i].text = fmt(v)
        shade(cells[i], rgba)
        if is_dark(rgba):
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    cells[5].text = f"{r['total_disp']:.2f}"

for row in table.rows:
    for c, w in zip(row.cells, [1.45, 0.95, 1.05, 0.95, 1.05, 0.95]):
        c.width = Inches(w)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for j, c in enumerate(row.cells):
        for p in c.paragraphs:
            if j:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.5)
for i in (0, 1):
    for c in table.rows[i].cells:
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True

key = doc.add_paragraph()
key.add_run("Shading: ").bold = True
key.add_run("blue = expert coders scored higher · neutral = the two "
            "codings agree · yellow = the LLM scored higher. "
            "Clamped at ±0.55.").italic = True
for run in key.runs:
    run.font.size = Pt(8.5)

doc.save(OUT / "table4_shifts.docx")
print("wrote", OUT / "table4_shifts.md")
print("wrote", OUT / "table4_shifts.docx")
