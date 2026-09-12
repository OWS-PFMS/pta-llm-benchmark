#!/usr/bin/env python3
"""Table 3 as a single landscape-page Word table (per Jeff, 2026-08-16):
challenge comparison on the left half, solution comparison on the right.

Built with python-docx (pandoc can't mix page orientations or merge header
cells). Data comes straight from analysis/loops_*.csv, same compaction rules
as make_tables.py.
"""
import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "analysis"
OUT = Path(__file__).resolve().parent

COMPACT = {"Community / Users": "Community/Users",
           "Govt / External Finance": "Govt/Ext Finance",
           "Laws + Regulations": "Laws + Regs"}

def top_loops(name, n=10):
    with open(DATA / f"loops_{name}.csv") as f:
        rows = list(csv.DictReader(f))
    out = []
    for r in rows[:n]:
        loop = r["loop"]
        for k, v in COMPACT.items():
            loop = loop.replace(k, v)
        parts = loop.split(" -> ")
        if len(parts) > 1 and parts[0] == parts[-1]:
            parts = parts[:-1]
        out.append((" → ".join(parts) + " ↺", float(r["score"])))
    return out

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Inches(0.6))

cap = doc.add_paragraph()
r = cap.add_run("Table 3. ")
r.bold = True
cap.add_run(
    "Ten highest-scoring feedback loops in each network, by coder and "
    "system. Loops are scored with the salience metric of Gottschamer and "
    "Walters (2023); ↺ marks the closure of the cycle back to its first "
    "factor.")

table = doc.add_table(rows=2, cols=9)
table.style = "Table Grid"

# Header: Rank spans both header rows; Challenge/Solution each span 4 cols.
table.cell(0, 0).merge(table.cell(1, 0)).text = "Rank"
table.cell(0, 1).merge(table.cell(0, 4)).text = "Challenge system"
table.cell(0, 5).merge(table.cell(0, 8)).text = "Solution system"
sub = ["Expert (benchmark)", "Score", "LLM (Claude)", "Score"]
for i, t in enumerate(sub * 2, start=1):
    table.cell(1, i).text = t

ch_nj, ch_cl = top_loops("nj_challenge"), top_loops("claude_challenge")
so_nj, so_cl = top_loops("nj_solution"), top_loops("claude_solution")
for i in range(10):
    row = table.add_row().cells
    cells = [str(i + 1),
             ch_nj[i][0], f"{ch_nj[i][1]:.1f}", ch_cl[i][0],
             f"{ch_cl[i][1]:.1f}",
             so_nj[i][0], f"{so_nj[i][1]:.1f}", so_cl[i][0],
             f"{so_cl[i][1]:.1f}"]
    for c, t in zip(row, cells):
        c.text = t

widths = [0.4, 1.85, 0.5, 1.85, 0.5, 1.85, 0.5, 1.85, 0.5]
for row in table.rows:
    for c, w in zip(row.cells, widths):
        c.width = Inches(w)
for row in table.rows:
    for c in row.cells:
        for p in c.paragraphs:
            for run in p.runs:
                run.font.size = Pt(8)
            if not p.runs:
                p.add_run("")
for i in (0, 1):  # bold headers
    for c in table.rows[i].cells:
        for p in c.paragraphs:
            for run in p.runs:
                run.bold = True

doc.save(OUT / "table3_landscape.docx")
print("wrote", OUT / "table3_landscape.docx")
